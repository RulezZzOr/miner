import tempfile
import subprocess
import sys
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from plugins.tools._reader_core import _guard_office_archive
from plugins.tools._reader_xlsx import _x_chart_lines
from scripts.security_scan import scan_files, scan_language


class DeliverySecurityTests(unittest.TestCase):
    def test_chart_xml_rejects_entity_declarations(self):
        # The chart reader itself must refuse a DTD. The standard library parser
        # would expand the entity and silently yield no chart, so this fails if the
        # reader stops using defusedxml with forbid_dtd=True.
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chart.xlsx"
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/charts/chart1.xml",
                                 '<!DOCTYPE root [<!ENTITY internal "expanded">]><root>&internal;</root>')
            self.assertEqual(_x_chart_lines(path), {"": ["Chart omitted: invalid or unsafe XML"]})

    def test_oversized_chart_is_not_decompressed(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chart.xlsx"
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/charts/chart1.xml", "x" * 2_000_001)
            with patch.object(zipfile.ZipFile, "read", side_effect=AssertionError("must not decompress")):
                result = _x_chart_lines(path)
            self.assertIn("safety limit", str(result))

    def test_office_zip_bomb_is_rejected_before_parsing(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "bomb.xlsx"
            with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
                archive.writestr("xl/sharedStrings.xml", b"x" * 25_000_001)
            with self.assertRaisesRegex(ValueError, "25 MB"):
                _guard_office_archive(path)
            from plugins.tools._doc_reader import reader_src
            result = subprocess.run([sys.executable, "-", str(path)], input=reader_src(), text=True, capture_output=True, timeout=10)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(result.stdout, "")
            self.assertIn("Office archive rejected", result.stderr)

    def test_small_real_spreadsheet_still_loads(self):
        from openpyxl import Workbook, load_workbook
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "safe.xlsx"
            book = Workbook(); book.active["A1"] = "OK"; book.save(path); book.close()
            _guard_office_archive(path)
            loaded = load_workbook(path)
            self.assertEqual(loaded.active["A1"].value, "OK")
            loaded.close()

    def test_publication_scan_detects_credentials_without_printing_values(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            value = "gh" + "p_" + "a" * 36
            (root / "config.txt").write_text(value)
            (root / ".ENV").write_text("local configuration")
            (root / "alias").symlink_to(root / "config.txt")
            findings = scan_files(root, ["config.txt", ".ENV", "alias"])
            self.assertEqual({f["kind"] for f in findings}, {"GitHub token", "private runtime path", "tracked symlink"})
            self.assertNotIn(value, str(findings))

    def test_publication_scan_flags_key_files_by_name_without_reading_them(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            names = ["miner-access-key.txt", "server.key", "bundle.PEM", "client.p12",
                     "analysis/audit-2026/remediation/resume-state.json", "README.md"]
            for name in names:
                (root / name).parent.mkdir(parents=True, exist_ok=True)
                (root / name).write_text("token-shaped-" + "x" * 30)
            findings = scan_files(root, names)
            self.assertEqual({f["path"] for f in findings}, set(names) - {"README.md"})
            self.assertEqual({f["kind"] for f in findings}, {"private runtime path"})

    def test_language_scan_rejects_czech_but_keeps_other_accents_and_upstream_data(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            # Escapes keep the probe words out of this file's own source text.
            samples = {"notes.md": "Start selh\x61l.\n", "ui.js": "label = 'P\u0159ehled'\n",
                       "ok.py": "slug('Caf\u00e9 Z\u00fcrich') == 'cafe-zurich'\n",
                       "frontier/benchmarks/data.txt": "P\u0159ehled\n"}
            for name, text in samples.items():
                (root / name).parent.mkdir(parents=True, exist_ok=True)
                (root / name).write_text(text, encoding="utf-8")
            findings = scan_language(root, list(samples))
            self.assertEqual({f["path"] for f in findings}, {"notes.md", "ui.js"})
            self.assertTrue(all(f["line"] == 1 for f in findings))

    def test_reader_bundle_reads_text_without_optional_xml_dependency(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "plain.txt"
            path.write_text("hello world\n")
            from plugins.tools._doc_reader import reader_src
            # -S drops site-packages, so defusedxml is unavailable to the bundle.
            result = subprocess.run([sys.executable, "-S", "-", str(path), "8000"], input=reader_src(),
                                    text=True, capture_output=True, timeout=30)
            self.assertEqual(result.returncode, 0, result.stderr[-2000:])
            self.assertIn("hello world", result.stdout)

    def test_chart_summary_degrades_when_defusedxml_is_missing(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chart.xlsx"
            with zipfile.ZipFile(path, "w") as archive:
                archive.writestr("xl/charts/chart1.xml", "<root/>")
            with patch.dict(sys.modules, {"defusedxml": None, "defusedxml.ElementTree": None}):
                result = _x_chart_lines(path)
            self.assertIn("defusedxml unavailable", str(result))
